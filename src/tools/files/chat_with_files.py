"""
Chat With Files Manager for AI Companion
Enables intelligent conversation with file contents using embeddings and RAG
"""

import os
import json
import sqlite3
import hashlib
from typing import Dict, Any, List, Optional, Tuple
from datetime import datetime
import numpy as np
from pathlib import Path
import mimetypes
import re

# Document processing
try:
    import PyPDF2
    import pdfplumber
    HAS_PDF = True
except ImportError:
    HAS_PDF = False

try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import pandas as pd
    HAS_PANDAS = True
except ImportError:
    HAS_PANDAS = False

# Embeddings
try:
    from sentence_transformers import SentenceTransformer
    HAS_SENTENCE_TRANSFORMERS = True
except ImportError:
    HAS_SENTENCE_TRANSFORMERS = False

try:
    import openai
    HAS_OPENAI = True
except ImportError:
    HAS_OPENAI = False

class FileChunk:
    """Represents a chunk of a file with metadata"""
    def __init__(self, content: str, file_path: str, chunk_id: int, 
                 start_pos: int = 0, end_pos: int = 0, metadata: Dict = None):
        self.content = content
        self.file_path = file_path
        self.chunk_id = chunk_id
        self.start_pos = start_pos
        self.end_pos = end_pos
        self.metadata = metadata or {}
        self.embedding = None
        self.content_hash = hashlib.md5(content.encode()).hexdigest()

class ChatWithFilesManager:
    """Manages file indexing and chat functionality"""
    
    def __init__(self, data_dir: str = "data/chat_files"):
        """Initialize chat with files manager"""
        self.data_dir = Path(data_dir)
        self.data_dir.mkdir(parents=True, exist_ok=True)
        
        # Database for storing file chunks and embeddings
        self.db_path = self.data_dir / "files.db"
        self.init_database()
        
        # Embedding model
        self.embedding_model = None
        self.embedding_dim = 384  # default for all-MiniLM-L6-v2
        self._initialize_embeddings()
        
        # File processors
        self.processors = {
            '.txt': self._process_text_file,
            '.md': self._process_text_file,
            '.py': self._process_code_file,
            '.js': self._process_code_file,
            '.html': self._process_code_file,
            '.css': self._process_code_file,
            '.json': self._process_json_file,
            '.csv': self._process_csv_file,
            '.pdf': self._process_pdf_file,
            '.docx': self._process_docx_file,
        }
        
        # Chunk settings
        self.chunk_size = 500  # characters per chunk
        self.chunk_overlap = 50  # overlap between chunks
        
        print(f"[ChatFiles] Initialized with {len(self.get_indexed_files())} indexed files")
    
    def init_database(self):
        """Initialize SQLite database for file storage"""
        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()
            
            # Files table
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS files (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    file_path TEXT UNIQUE NOT NULL,
                    file_name TEXT NOT NULL,
                    file_size INTEGER,
                    file_type TEXT,
                    content_hash TEXT,
                    indexed_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    last_modified TIMESTAMP,
                    chunk_count INTEGER DEFAULT 0
                )
            ''')
            
            # Chunks table
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS chunks (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    file_id INTEGER,
                    chunk_id INTEGER,
                    content TEXT NOT NULL,
                    content_hash TEXT,
                    start_pos INTEGER,
                    end_pos INTEGER,
                    metadata TEXT,
                    embedding BLOB,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY (file_id) REFERENCES files (id)
                )
            ''')
            
            # Create indexes
            cursor.execute('CREATE INDEX IF NOT EXISTS idx_file_path ON files (file_path)')
            cursor.execute('CREATE INDEX IF NOT EXISTS idx_chunk_file ON chunks (file_id)')
            cursor.execute('CREATE INDEX IF NOT EXISTS idx_content_hash ON chunks (content_hash)')
            
            conn.commit()
    
    def _initialize_embeddings(self):
        """Initialize embedding model"""
        openai_key = os.getenv("OPENAI_API_KEY")
        
        if HAS_SENTENCE_TRANSFORMERS:
            try:
                self.embedding_model = SentenceTransformer('all-MiniLM-L6-v2')
                self.embedding_method = "sentence_transformers"
                self.embedding_dim = 384
                print("[ChatFiles] Using SentenceTransformers for embeddings")
                return
            except Exception as e:
                print(f"[ChatFiles] SentenceTransformers failed: {e}")
        
        if HAS_OPENAI and openai_key:
            try:
                openai.api_key = openai_key
                self.embedding_method = "openai"
                self.embedding_dim = 1536  # text-embedding-ada-002 dimension
                print("[ChatFiles] Using OpenAI for embeddings")
                return
            except Exception as e:
                print(f"[ChatFiles] OpenAI embeddings failed: {e}")
        
        print("[ChatFiles] ⚠️  No embedding model available - search will be limited")
        self.embedding_method = "none"
    
    def _get_file_embedding(self, text: str) -> Optional[np.ndarray]:
        """Generate embedding for text"""
        if self.embedding_method == "sentence_transformers" and self.embedding_model:
            try:
                return self.embedding_model.encode(text)
            except Exception as e:
                print(f"[ChatFiles] Embedding error: {e}")
                return None
        
        elif self.embedding_method == "openai":
            try:
                response = openai.Embedding.create(
                    input=text,
                    model="text-embedding-ada-002"
                )
                return np.array(response['data'][0]['embedding'])
            except Exception as e:
                print(f"[ChatFiles] OpenAI embedding error: {e}")
                return None
        
        return None
    
    def add_file(self, file_path: str, force_reindex: bool = False) -> Dict[str, Any]:
        """Add a file to the chat system"""
        file_path = str(Path(file_path).resolve())
        
        if not os.path.exists(file_path):
            return {"success": False, "error": f"File not found: {file_path}"}
        
        file_name = os.path.basename(file_path)
        file_size = os.path.getsize(file_path)
        file_type = mimetypes.guess_type(file_path)[0] or "unknown"
        last_modified = datetime.fromtimestamp(os.path.getmtime(file_path))
        
        # Check if already indexed
        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()
            cursor.execute('SELECT id, content_hash FROM files WHERE file_path = ?', (file_path,))
            existing = cursor.fetchone()
        
        # Calculate content hash
        content_hash = self._calculate_file_hash(file_path)
        
        if existing and not force_reindex:
            file_id, old_hash = existing
            if old_hash == content_hash:
                return {"success": True, "message": "File already indexed", "file_id": file_id}
        
        # Process the file
        try:
            chunks = self._process_file(file_path)
            if not chunks:
                return {"success": False, "error": "Could not process file - unsupported format"}
            
            # Store in database
            with sqlite3.connect(self.db_path) as conn:
                cursor = conn.cursor()
                
                if existing:
                    # Update existing file
                    file_id = existing[0]
                    cursor.execute('''
                        UPDATE files SET file_size = ?, content_hash = ?, 
                        indexed_at = CURRENT_TIMESTAMP, last_modified = ?, chunk_count = ?
                        WHERE id = ?
                    ''', (file_size, content_hash, last_modified, len(chunks), file_id))
                    
                    # Delete old chunks
                    cursor.execute('DELETE FROM chunks WHERE file_id = ?', (file_id,))
                else:
                    # Insert new file
                    cursor.execute('''
                        INSERT INTO files (file_path, file_name, file_size, file_type, 
                        content_hash, last_modified, chunk_count)
                        VALUES (?, ?, ?, ?, ?, ?, ?)
                    ''', (file_path, file_name, file_size, file_type, content_hash, 
                          last_modified, len(chunks)))
                    file_id = cursor.lastrowid
                
                # Insert chunks
                for chunk in chunks:
                    embedding_blob = None
                    if chunk.embedding is not None:
                        embedding_blob = chunk.embedding.tobytes()
                    
                    cursor.execute('''
                        INSERT INTO chunks (file_id, chunk_id, content, content_hash,
                        start_pos, end_pos, metadata, embedding)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (file_id, chunk.chunk_id, chunk.content, chunk.content_hash,
                          chunk.start_pos, chunk.end_pos, json.dumps(chunk.metadata),
                          embedding_blob))
                
                conn.commit()
            
            return {
                "success": True,
                "message": f"File indexed with {len(chunks)} chunks",
                "file_id": file_id,
                "chunks": len(chunks)
            }
        
        except Exception as e:
            return {"success": False, "error": f"Failed to process file: {str(e)}"}
    
    def _calculate_file_hash(self, file_path: str) -> str:
        """Calculate MD5 hash of file content"""
        hash_md5 = hashlib.md5()
        with open(file_path, "rb") as f:
            for chunk in iter(lambda: f.read(4096), b""):
                hash_md5.update(chunk)
        return hash_md5.hexdigest()
    
    def _process_file(self, file_path: str) -> List[FileChunk]:
        """Process a file into chunks"""
        file_ext = Path(file_path).suffix.lower()
        
        if file_ext in self.processors:
            try:
                return self.processors[file_ext](file_path)
            except Exception as e:
                print(f"[ChatFiles] Error processing {file_path}: {e}")
                # Fallback to text processing
                return self._process_text_file(file_path)
        else:
            # Try as text file
            return self._process_text_file(file_path)
    
    def _process_text_file(self, file_path: str) -> List[FileChunk]:
        """Process plain text file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
        except UnicodeDecodeError:
            with open(file_path, 'r', encoding='latin-1') as f:
                content = f.read()
        
        return self._create_text_chunks(content, file_path)
    
    def _process_code_file(self, file_path: str) -> List[FileChunk]:
        """Process code file with syntax awareness"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
        except UnicodeDecodeError:
            return []
        
        # For code files, try to chunk by functions/classes
        chunks = []
        lines = content.split('\n')
        current_chunk = ""
        chunk_id = 0
        start_line = 0
        
        # Simple code structure detection
        function_pattern = r'^(def |function |class |public |private |\w+\s*\(.*\)\s*\{)'
        
        for i, line in enumerate(lines):
            current_chunk += line + '\n'
            
            # If we hit a function/class definition and have enough content
            if (re.match(function_pattern, line.strip()) and len(current_chunk) > 200) or \
               len(current_chunk) > self.chunk_size:
                
                if current_chunk.strip():
                    embedding = self._get_file_embedding(current_chunk)
                    chunk = FileChunk(
                        content=current_chunk,
                        file_path=file_path,
                        chunk_id=chunk_id,
                        start_pos=start_line,
                        end_pos=i,
                        metadata={"type": "code", "language": Path(file_path).suffix[1:]}
                    )
                    chunk.embedding = embedding
                    chunks.append(chunk)
                
                current_chunk = line + '\n'
                start_line = i
                chunk_id += 1
        
        # Add remaining content
        if current_chunk.strip():
            embedding = self._get_file_embedding(current_chunk)
            chunk = FileChunk(
                content=current_chunk,
                file_path=file_path,
                chunk_id=chunk_id,
                start_pos=start_line,
                end_pos=len(lines),
                metadata={"type": "code", "language": Path(file_path).suffix[1:]}
            )
            chunk.embedding = embedding
            chunks.append(chunk)
        
        return chunks
    
    def _process_json_file(self, file_path: str) -> List[FileChunk]:
        """Process JSON file"""
        with open(file_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        # Convert JSON to readable text
        content = json.dumps(data, indent=2)
        return self._create_text_chunks(content, file_path, metadata={"type": "json"})
    
    def _process_csv_file(self, file_path: str) -> List[FileChunk]:
        """Process CSV file"""
        if not HAS_PANDAS:
            return self._process_text_file(file_path)
        
        try:
            df = pd.read_csv(file_path)
            
            # Create summary chunk
            summary = f"CSV File: {Path(file_path).name}\n"
            summary += f"Rows: {len(df)}, Columns: {len(df.columns)}\n"
            summary += f"Columns: {', '.join(df.columns.tolist())}\n\n"
            summary += "Sample data:\n"
            summary += df.head(10).to_string()
            
            chunks = []
            
            # Summary chunk
            embedding = self._get_file_embedding(summary)
            chunk = FileChunk(
                content=summary,
                file_path=file_path,
                chunk_id=0,
                metadata={"type": "csv_summary", "rows": len(df), "columns": list(df.columns)}
            )
            chunk.embedding = embedding
            chunks.append(chunk)
            
            # Data chunks (every 100 rows)
            chunk_id = 1
            for i in range(0, len(df), 100):
                chunk_data = df.iloc[i:i+100]
                content = f"Rows {i}-{min(i+99, len(df)-1)}:\n"
                content += chunk_data.to_string(index=False)
                
                embedding = self._get_file_embedding(content)
                chunk = FileChunk(
                    content=content,
                    file_path=file_path,
                    chunk_id=chunk_id,
                    metadata={"type": "csv_data", "start_row": i, "end_row": min(i+99, len(df)-1)}
                )
                chunk.embedding = embedding
                chunks.append(chunk)
                chunk_id += 1
            
            return chunks
        
        except Exception as e:
            print(f"[ChatFiles] CSV processing failed: {e}")
            return self._process_text_file(file_path)
    
    def _process_pdf_file(self, file_path: str) -> List[FileChunk]:
        """Process PDF file"""
        if not HAS_PDF:
            return []
        
        text_content = ""
        
        try:
            # Try pdfplumber first (better text extraction)
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                for page_num, page in enumerate(pdf.pages):
                    page_text = page.extract_text()
                    if page_text:
                        text_content += f"\n--- Page {page_num + 1} ---\n"
                        text_content += page_text
        except:
            try:
                # Fallback to PyPDF2
                with open(file_path, 'rb') as f:
                    pdf_reader = PyPDF2.PdfReader(f)
                    for page_num, page in enumerate(pdf_reader.pages):
                        page_text = page.extract_text()
                        if page_text:
                            text_content += f"\n--- Page {page_num + 1} ---\n"
                            text_content += page_text
            except Exception as e:
                print(f"[ChatFiles] PDF processing failed: {e}")
                return []
        
        if not text_content.strip():
            return []
        
        return self._create_text_chunks(text_content, file_path, metadata={"type": "pdf"})
    
    def _process_docx_file(self, file_path: str) -> List[FileChunk]:
        """Process DOCX file"""
        if not HAS_DOCX:
            return []
        
        try:
            doc = DocxDocument(file_path)
            content = ""
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    content += paragraph.text + "\n"
            
            # Add tables
            for table in doc.tables:
                content += "\n[TABLE]\n"
                for row in table.rows:
                    row_text = " | ".join([cell.text for cell in row.cells])
                    content += row_text + "\n"
                content += "[/TABLE]\n"
            
            return self._create_text_chunks(content, file_path, metadata={"type": "docx"})
        
        except Exception as e:
            print(f"[ChatFiles] DOCX processing failed: {e}")
            return []
    
    def _create_text_chunks(self, content: str, file_path: str, 
                           metadata: Dict = None) -> List[FileChunk]:
        """Create text chunks from content"""
        chunks = []
        chunk_id = 0
        
        # Split into sentences for better chunk boundaries
        sentences = re.split(r'[.!?]+', content)
        
        current_chunk = ""
        start_pos = 0
        
        for sentence in sentences:
            sentence = sentence.strip()
            if not sentence:
                continue
            
            # If adding this sentence would exceed chunk size, create a chunk
            if len(current_chunk) + len(sentence) > self.chunk_size and current_chunk:
                end_pos = start_pos + len(current_chunk)
                
                embedding = self._get_file_embedding(current_chunk)
                chunk = FileChunk(
                    content=current_chunk.strip(),
                    file_path=file_path,
                    chunk_id=chunk_id,
                    start_pos=start_pos,
                    end_pos=end_pos,
                    metadata=metadata or {}
                )
                chunk.embedding = embedding
                chunks.append(chunk)
                
                # Start new chunk with overlap
                overlap_start = max(0, len(current_chunk) - self.chunk_overlap)
                current_chunk = current_chunk[overlap_start:] + " " + sentence
                start_pos = end_pos - (len(current_chunk) - len(sentence) - 1)
                chunk_id += 1
            else:
                current_chunk += " " + sentence if current_chunk else sentence
        
        # Add final chunk
        if current_chunk.strip():
            embedding = self._get_file_embedding(current_chunk)
            chunk = FileChunk(
                content=current_chunk.strip(),
                file_path=file_path,
                chunk_id=chunk_id,
                start_pos=start_pos,
                end_pos=start_pos + len(current_chunk),
                metadata=metadata or {}
            )
            chunk.embedding = embedding
            chunks.append(chunk)
        
        return chunks
    
    def search_files(self, query: str, max_results: int = 10, 
                    file_paths: List[str] = None) -> List[Dict[str, Any]]:
        """Search through indexed files"""
        results = []
        
        # Get query embedding
        query_embedding = self._get_file_embedding(query)
        
        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()
            
            # Build query
            if file_paths:
                placeholders = ','.join(['?' for _ in file_paths])
                cursor.execute(f'''
                    SELECT c.content, c.metadata, c.embedding, f.file_path, f.file_name,
                           c.start_pos, c.end_pos, c.chunk_id
                    FROM chunks c
                    JOIN files f ON c.file_id = f.id
                    WHERE f.file_path IN ({placeholders})
                ''', file_paths)
            else:
                cursor.execute('''
                    SELECT c.content, c.metadata, c.embedding, f.file_path, f.file_name,
                           c.start_pos, c.end_pos, c.chunk_id
                    FROM chunks c
                    JOIN files f ON c.file_id = f.id
                ''')
            
            rows = cursor.fetchall()
        
        # Calculate similarities and text relevance
        for row in rows:
            content, metadata_json, embedding_blob, file_path, file_name, start_pos, end_pos, chunk_id = row
            
            similarity = 0.0
            
            # Vector similarity if embeddings available
            if query_embedding is not None and embedding_blob:
                try:
                    chunk_embedding = np.frombuffer(embedding_blob, dtype=np.float32)
                    if len(chunk_embedding) == len(query_embedding):
                        similarity = np.dot(query_embedding, chunk_embedding) / (
                            np.linalg.norm(query_embedding) * np.linalg.norm(chunk_embedding)
                        )
                except Exception as e:
                    print(f"[ChatFiles] Embedding similarity error: {e}")
            
            # Text-based relevance (keyword matching)
            query_lower = query.lower()
            content_lower = content.lower()
            text_score = 0.0
            
            query_words = query_lower.split()
            for word in query_words:
                if word in content_lower:
                    text_score += content_lower.count(word) / len(content.split())
            
            # Combined score
            combined_score = max(similarity, text_score * 0.5)  # Weight vector similarity more
            
            if combined_score > 0.1:  # Minimum relevance threshold
                try:
                    metadata = json.loads(metadata_json) if metadata_json else {}
                except:
                    metadata = {}
                
                results.append({
                    "content": content,
                    "file_path": file_path,
                    "file_name": file_name,
                    "chunk_id": chunk_id,
                    "start_pos": start_pos,
                    "end_pos": end_pos,
                    "metadata": metadata,
                    "similarity": similarity,
                    "text_score": text_score,
                    "combined_score": combined_score
                })
        
        # Sort by combined score and return top results
        results.sort(key=lambda x: x["combined_score"], reverse=True)
        return results[:max_results]
    
    def get_file_content(self, file_path: str, chunk_id: int = None) -> Dict[str, Any]:
        """Get content of a specific file or chunk"""
        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()
            
            if chunk_id is not None:
                cursor.execute('''
                    SELECT c.content, c.metadata, f.file_name
                    FROM chunks c
                    JOIN files f ON c.file_id = f.id
                    WHERE f.file_path = ? AND c.chunk_id = ?
                ''', (file_path, chunk_id))
                row = cursor.fetchone()
                
                if row:
                    content, metadata_json, file_name = row
                    metadata = json.loads(metadata_json) if metadata_json else {}
                    return {
                        "success": True,
                        "content": content,
                        "file_name": file_name,
                        "chunk_id": chunk_id,
                        "metadata": metadata
                    }
            else:
                cursor.execute('''
                    SELECT c.content, c.chunk_id, c.metadata
                    FROM chunks c
                    JOIN files f ON c.file_id = f.id
                    WHERE f.file_path = ?
                    ORDER BY c.chunk_id
                ''', (file_path,))
                rows = cursor.fetchall()
                
                if rows:
                    full_content = ""
                    chunks = []
                    for content, chunk_id, metadata_json in rows:
                        full_content += content + "\n\n"
                        metadata = json.loads(metadata_json) if metadata_json else {}
                        chunks.append({
                            "chunk_id": chunk_id,
                            "content": content,
                            "metadata": metadata
                        })
                    
                    return {
                        "success": True,
                        "content": full_content.strip(),
                        "file_name": os.path.basename(file_path),
                        "chunks": chunks
                    }
        
        return {"success": False, "error": "File or chunk not found"}
    
    def get_indexed_files(self) -> List[Dict[str, Any]]:
        """Get list of all indexed files"""
        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()
            cursor.execute('''
                SELECT file_path, file_name, file_size, file_type, 
                       indexed_at, chunk_count, content_hash
                FROM files
                ORDER BY indexed_at DESC
            ''')
            
            files = []
            for row in cursor.fetchall():
                file_path, file_name, file_size, file_type, indexed_at, chunk_count, content_hash = row
                files.append({
                    "file_path": file_path,
                    "file_name": file_name,
                    "file_size": file_size,
                    "file_type": file_type,
                    "indexed_at": indexed_at,
                    "chunk_count": chunk_count,
                    "content_hash": content_hash
                })
            
            return files
    
    def remove_file(self, file_path: str) -> Dict[str, Any]:
        """Remove a file from the index"""
        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()
            
            # Get file ID
            cursor.execute('SELECT id FROM files WHERE file_path = ?', (file_path,))
            row = cursor.fetchone()
            
            if not row:
                return {"success": False, "error": "File not found in index"}
            
            file_id = row[0]
            
            # Delete chunks first (foreign key constraint)
            cursor.execute('DELETE FROM chunks WHERE file_id = ?', (file_id,))
            chunks_deleted = cursor.rowcount
            
            # Delete file record
            cursor.execute('DELETE FROM files WHERE id = ?', (file_id,))
            
            conn.commit()
            
            return {
                "success": True,
                "message": f"Removed file and {chunks_deleted} chunks from index"
            }
    
    def chat_with_files(self, query: str, context_files: List[str] = None,
                       max_context_chunks: int = 5) -> Dict[str, Any]:
        """Generate response based on file contents"""
        
        # Search for relevant content
        search_results = self.search_files(query, max_results=max_context_chunks, 
                                         file_paths=context_files)
        
        if not search_results:
            return {
                "success": False,
                "error": "No relevant content found in indexed files"
            }
        
        # Build context from search results
        context = "Based on the following file contents:\n\n"
        
        for i, result in enumerate(search_results):
            context += f"--- From {result['file_name']} (chunk {result['chunk_id']}) ---\n"
            context += result['content']
            context += f"\n(Relevance: {result['combined_score']:.2f})\n\n"
        
        context += f"Question: {query}\n\n"
        context += "Answer based on the provided file contents:"
        
        return {
            "success": True,
            "context": context,
            "sources": [
                {
                    "file_name": r['file_name'],
                    "file_path": r['file_path'],
                    "chunk_id": r['chunk_id'],
                    "relevance": r['combined_score']
                } for r in search_results
            ],
            "query": query
        }

# Demo function
def demo_chat_with_files():
    """Demo chat with files functionality"""
    print("💬 Chat With Files Demo")
    print("=" * 30)
    
    chat_manager = ChatWithFilesManager()
    
    # Create a sample file
    sample_file = Path("sample_document.txt")
    sample_content = """
    Artificial Intelligence and Machine Learning
    
    Machine learning is a subset of artificial intelligence (AI) that focuses on the development of algorithms and statistical models that enable computer systems to perform specific tasks without using explicit instructions, relying on patterns and inference instead.
    
    Deep Learning
    Deep learning is part of a broader family of machine learning methods based on artificial neural networks with representation learning. Learning can be supervised, semi-supervised or unsupervised.
    
    Natural Language Processing
    Natural language processing (NLP) is a subfield of linguistics, computer science, and artificial intelligence concerned with the interactions between computers and human language, in particular how to program computers to process and analyze large amounts of natural language data.
    """
    
    print("📄 Creating sample document...")
    with open(sample_file, 'w') as f:
        f.write(sample_content)
    
    # Index the file
    print("🔍 Indexing sample document...")
    result = chat_manager.add_file(str(sample_file))
    print(f"   {result}")
    
    # Show indexed files
    print("\n📋 Indexed files:")
    files = chat_manager.get_indexed_files()
    for file_info in files:
        print(f"   • {file_info['file_name']} ({file_info['chunk_count']} chunks)")
    
    # Test search
    test_queries = [
        "What is machine learning?",
        "Tell me about deep learning",
        "How does NLP work?",
        "artificial intelligence applications"
    ]
    
    print("\n🔍 Testing search:")
    for query in test_queries:
        results = chat_manager.search_files(query, max_results=3)
        print(f"\n   Query: '{query}'")
        if results:
            for result in results:
                print(f"   → Score: {result['combined_score']:.2f}")
                print(f"     Content: {result['content'][:100]}...")
        else:
            print("   → No results found")
    
    # Test chat functionality
    print("\n💬 Testing chat with files:")
    chat_result = chat_manager.chat_with_files("Explain the relationship between AI and machine learning")
    if chat_result["success"]:
        print("✅ Context generated successfully!")
        print(f"   Sources: {len(chat_result['sources'])}")
        print(f"   Context length: {len(chat_result['context'])} characters")
    else:
        print(f"❌ Error: {chat_result['error']}")
    
    # Cleanup
    if sample_file.exists():
        sample_file.unlink()
        print("\n🧹 Cleaned up sample file")

if __name__ == "__main__":
    demo_chat_with_files()
